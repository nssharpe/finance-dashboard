"""Historical xlsx importer for the local personal-finance dashboard.

Reads a historical spending workbook (read-only) and populates:
  - ``transactions`` table — Mint 2018-2023, Empower 2024, CSV 2025
  - ``annual_totals`` table — Summary sheet categories for 2013-2017

Run as a module::

    python -m app.importers.xlsx_history

or call ``import_all()`` programmatically.
"""

from __future__ import annotations

import json
import os
from collections import defaultdict
from datetime import datetime
from typing import TYPE_CHECKING

import openpyxl

from app import db
from app.categorize import categorize, map_source_category

if TYPE_CHECKING:
    import sqlite3

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

_HERE = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.dirname(os.path.dirname(_HERE))

def _find_workbook() -> str:
    """Locate the history workbook.

    ``FINANCE_HISTORY_XLSX`` wins if set (absolute, or relative to the project
    root). Otherwise fall back to the single .xlsx in the project root, which is
    the common case — workbooks are gitignored, so the filename is personal and
    can't be hardcoded.
    """
    override = os.environ.get("FINANCE_HISTORY_XLSX")
    if override:
        return override if os.path.isabs(override) else os.path.join(_PROJECT_ROOT, override)

    candidates = sorted(
        f for f in os.listdir(_PROJECT_ROOT)
        if f.lower().endswith(".xlsx") and not f.startswith("~$")
    ) if os.path.isdir(_PROJECT_ROOT) else []
    if len(candidates) == 1:
        return os.path.join(_PROJECT_ROOT, candidates[0])
    # Ambiguous or absent: return a stable path so the caller's "file not found"
    # error names something sensible.
    return os.path.join(_PROJECT_ROOT, "Spending Analysis.xlsx")


DEFAULT_WORKBOOK_PATH = _find_workbook()

# ---------------------------------------------------------------------------
# Sheet configuration
# ---------------------------------------------------------------------------

# (sheet_name, source_tag, format)
# format: "mint" | "empower" | "csv2025"
DETAIL_SHEETS: list[tuple[str, str, str]] = [
    ("Mint Details 2018", "mint", "mint"),
    ("Mint Details 2019", "mint", "mint"),
    ("Mint Details 2020", "mint", "mint"),
    ("Mint Details 2021", "mint", "mint"),
    ("Mint Details 2022", "mint", "mint"),
    ("Mint Details 2023", "mint", "mint"),
    ("Empower Details 2024", "empower", "empower"),
    ("2025 Consolidated Spending", "csv2025", "csv2025"),
]

# Only import annual_totals for years without transaction detail.
ANNUAL_TOTAL_YEARS = {2013, 2014, 2015, 2016, 2017}

# Summary row categories to skip entirely (whole-portfolio aggregates).
_SKIP_CATEGORY_PREFIXES = ("Total",)


# ---------------------------------------------------------------------------
# Helper: robust ISO date from openpyxl cell value
# ---------------------------------------------------------------------------

def _to_iso(val) -> str | None:
    """Convert an openpyxl date/datetime cell value to 'YYYY-MM-DD', or None."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        val = val.strip()
        if not val:
            return None
        # Try a few common formats.
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
            try:
                return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


# ---------------------------------------------------------------------------
# Sheet importers
# ---------------------------------------------------------------------------

def _import_mint_sheet(
    conn: "sqlite3.Connection",
    ws,
    source: str,
    per_year: dict[int, int],
    unmapped: dict[str, int],
) -> int:
    """Import one Mint-format sheet (2018-2023).

    Columns: Date | Description | Original Description | Amount |
             Transaction Type | Category | Account Name | Labels | Notes

    Sign convention: Transaction Type is 'debit' (money out) or 'credit' (money in).
    Amount is the positive magnitude — we apply sign here.
    """
    imported = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        date_val, description, orig_description, amount, txn_type, category_raw, account_name = (
            row[0], row[1], row[2], row[3], row[4], row[5], row[6]
        )

        iso_date = _to_iso(date_val)
        if not iso_date or amount is None:
            continue

        try:
            amount = float(amount)
        except (TypeError, ValueError):
            continue

        # Apply sign: debit = money out = negative.
        txn_type_str = (txn_type or "").strip().lower()
        if txn_type_str == "debit":
            signed_amount = -abs(amount)
        elif txn_type_str == "credit":
            signed_amount = abs(amount)
        else:
            # Unknown type — keep as-is (shouldn't occur with clean Mint data).
            signed_amount = amount

        # Prefer Original Description when present.
        stored_desc = (orig_description or "").strip() or (description or "").strip()

        source_category = (category_raw or "").strip() or None

        # Track unmapped source categories.
        if source_category and map_source_category(source_category) is None:
            unmapped[source_category] += 1

        canon_category = categorize(stored_desc, source_category)

        account = (account_name or "").strip() or "Unknown"
        txn_id = db.synthetic_id(iso_date, account, stored_desc, signed_amount)

        db.upsert_transaction(
            conn,
            {
                "id": txn_id,
                "account_id": None,
                "account_name": account,
                "posted": iso_date,
                "amount": signed_amount,
                "description": stored_desc,
                "payee": None,
                "memo": None,
                "pending": 0,
                "source": source,
                "source_category": source_category,
                "category": canon_category,
            },
        )

        year = int(iso_date[:4])
        per_year[year] = per_year.get(year, 0) + 1
        imported += 1

    return imported


def _import_empower_sheet(
    conn: "sqlite3.Connection",
    ws,
    source: str,
    per_year: dict[int, int],
    unmapped: dict[str, int],
) -> int:
    """Import the Empower 2024 sheet.

    Columns: Date | Account | Description | Category | Tags | Amount

    Amount is already signed (negative = spending).
    """
    imported = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        date_val, account_name, description, category_raw, _tags, amount = (
            row[0], row[1], row[2], row[3], row[4], row[5]
        )

        iso_date = _to_iso(date_val)
        if not iso_date or amount is None:
            continue

        try:
            signed_amount = float(amount)
        except (TypeError, ValueError):
            continue

        stored_desc = (description or "").strip()
        source_category = (category_raw or "").strip() or None

        if source_category and map_source_category(source_category) is None:
            unmapped[source_category] += 1

        canon_category = categorize(stored_desc, source_category)

        account = (account_name or "").strip() or "Unknown"
        txn_id = db.synthetic_id(iso_date, account, stored_desc, signed_amount)

        db.upsert_transaction(
            conn,
            {
                "id": txn_id,
                "account_id": None,
                "account_name": account,
                "posted": iso_date,
                "amount": signed_amount,
                "description": stored_desc,
                "payee": None,
                "memo": None,
                "pending": 0,
                "source": source,
                "source_category": source_category,
                "category": canon_category,
            },
        )

        year = int(iso_date[:4])
        per_year[year] = per_year.get(year, 0) + 1
        imported += 1

    return imported


def _import_csv2025_sheet(
    conn: "sqlite3.Connection",
    ws,
    source: str,
    per_year: dict[int, int],
    unmapped: dict[str, int],
) -> int:
    """Import the 2025 Consolidated Spending sheet.

    Columns: Date | Account | Description | Amount | Category

    Amount is already signed (negative = spending).
    """
    imported = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        date_val, account_name, description, amount, category_raw = (
            row[0], row[1], row[2], row[3], row[4]
        )

        iso_date = _to_iso(date_val)
        if not iso_date or amount is None:
            continue

        try:
            signed_amount = float(amount)
        except (TypeError, ValueError):
            continue

        stored_desc = (description or "").strip()
        source_category = (category_raw or "").strip() or None

        if source_category and map_source_category(source_category) is None:
            unmapped[source_category] += 1

        canon_category = categorize(stored_desc, source_category)

        account = (account_name or "").strip() or "Unknown"
        txn_id = db.synthetic_id(iso_date, account, stored_desc, signed_amount)

        db.upsert_transaction(
            conn,
            {
                "id": txn_id,
                "account_id": None,
                "account_name": account,
                "posted": iso_date,
                "amount": signed_amount,
                "description": stored_desc,
                "payee": None,
                "memo": None,
                "pending": 0,
                "source": source,
                "source_category": source_category,
                "category": canon_category,
            },
        )

        year = int(iso_date[:4])
        per_year[year] = per_year.get(year, 0) + 1
        imported += 1

    return imported


# ---------------------------------------------------------------------------
# Annual totals (Summary sheet, 2013-2017 only)
# ---------------------------------------------------------------------------

def _import_annual_totals(conn: "sqlite3.Connection", ws) -> int:
    """Import Summary-sheet rows into annual_totals for ANNUAL_TOTAL_YEARS.

    Row 0 (index 0): first cell is blank/space, remaining cells are year floats.
    Row 1+: col 0 = category name, remaining = amounts aligned to year headers.
    Skip rows whose category is None/blank or starts with 'Total'.
    """
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return 0

    # Build year -> column-index mapping from header row.
    header = all_rows[0]
    year_col: dict[int, int] = {}
    for col_idx, cell_val in enumerate(header):
        if isinstance(cell_val, (int, float)):
            yr = int(cell_val)
            if yr in ANNUAL_TOTAL_YEARS:
                year_col[yr] = col_idx

    if not year_col:
        return 0

    imported = 0
    for row in all_rows[1:]:
        if not row:
            continue
        category = row[0]
        if not category or not str(category).strip():
            continue
        category = str(category).strip()

        # Skip aggregate rows.
        skip = any(category == prefix or category.startswith(prefix)
                   for prefix in _SKIP_CATEGORY_PREFIXES)
        if skip:
            continue

        for year, col_idx in year_col.items():
            if col_idx >= len(row):
                continue
            amount_val = row[col_idx]
            if amount_val is None:
                continue
            try:
                amount = float(amount_val)
            except (TypeError, ValueError):
                continue

            conn.execute(
                """
                INSERT INTO annual_totals (year, category, amount)
                VALUES (?, ?, ?)
                ON CONFLICT(year, category) DO UPDATE SET amount = excluded.amount
                """,
                (year, category, amount),
            )
            imported += 1

    return imported


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def import_all(
    db_path: str | None = None,
    workbook_path: str | None = None,
) -> dict:
    """Run the full historical import.

    Opens the DB (runs ``init_db``), imports all transaction-detail sheets and
    the 2013-2017 annual totals from the Summary sheet, calls
    ``db.apply_overrides``, and commits. Idempotent — synthetic ids dedupe on
    re-runs.

    Parameters
    ----------
    db_path:
        Path to the SQLite database file. Defaults to ``finance.db`` at the
        project root (same default as ``db.get_conn()``).
    workbook_path:
        Path to the source workbook. Defaults to the workbook that lives in
        the same directory as this project.

    Returns
    -------
    dict with keys:
        ``per_year``               — {year: transaction_count}
        ``transactions_total``     — total rows upserted
        ``annual_totals_rows``     — rows written to annual_totals
        ``unmapped_source_categories`` — {source_category: count} for categories
                                         that fell through to description rules
    """
    wb_path = workbook_path or DEFAULT_WORKBOOK_PATH
    conn = db.get_conn(db_path)
    db.init_db(conn)

    per_year: dict[int, int] = {}
    unmapped: dict[str, int] = defaultdict(int)
    transactions_total = 0

    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    try:
        for sheet_name, source_tag, fmt in DETAIL_SHEETS:
            if sheet_name not in wb.sheetnames:
                print(f"  [warn] sheet '{sheet_name}' not found, skipping.")
                continue

            ws = wb[sheet_name]
            print(f"  Importing '{sheet_name}' (format={fmt}) …", flush=True)

            if fmt == "mint":
                n = _import_mint_sheet(conn, ws, source_tag, per_year, unmapped)
            elif fmt == "empower":
                n = _import_empower_sheet(conn, ws, source_tag, per_year, unmapped)
            elif fmt == "csv2025":
                n = _import_csv2025_sheet(conn, ws, source_tag, per_year, unmapped)
            else:
                print(f"  [warn] unknown format '{fmt}' for '{sheet_name}', skipping.")
                n = 0

            print(f"    -> {n} rows", flush=True)
            transactions_total += n
            conn.commit()  # Commit after each sheet for resilience.

        # Annual totals from Summary sheet.
        print("  Importing annual_totals from 'Summary' …", flush=True)
        summary_ws = wb["Summary"]
        annual_rows = _import_annual_totals(conn, summary_ws)
        print(f"    -> {annual_rows} rows", flush=True)
        conn.commit()

    finally:
        wb.close()

    # Re-apply any manual category overrides.
    overrides_applied = db.apply_overrides(conn)
    conn.commit()
    print(f"  Overrides applied: {overrides_applied}", flush=True)

    conn.close()

    return {
        "per_year": dict(sorted(per_year.items())),
        "transactions_total": transactions_total,
        "annual_totals_rows": annual_rows,
        "unmapped_source_categories": dict(
            sorted(unmapped.items(), key=lambda kv: -kv[1])
        ),
    }


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sqlite3 as _sqlite3

    print("=" * 60)
    print("Historical xlsx importer")
    print("=" * 60)

    summary = import_all()

    print()
    print("=== Import summary ===")
    print(f"Transactions total : {summary['transactions_total']:,}")
    print(f"Annual totals rows : {summary['annual_totals_rows']:,}")
    print()
    print("Per-year transaction counts:")
    for yr, cnt in summary["per_year"].items():
        print(f"  {yr}: {cnt:,}")

    print()
    unmapped = summary["unmapped_source_categories"]
    if unmapped:
        print(f"Unmapped source categories ({len(unmapped)} distinct):")
        for cat, cnt in list(unmapped.items())[:30]:
            print(f"  {cnt:4d}  {cat}")
    else:
        print("Unmapped source categories: none — full coverage!")

    # -----------------------------------------------------------------------
    # Spot-check: 5 sample rows + 2023 category breakdown
    # -----------------------------------------------------------------------
    conn = db.get_conn()
    try:
        print()
        print("=== 5 sample imported rows ===")
        rows = conn.execute(
            """
            SELECT posted, account_name, description, amount, source_category, category
            FROM transactions
            WHERE source IN ('mint','empower','csv2025')
            ORDER BY RANDOM()
            LIMIT 5
            """
        ).fetchall()
        for r in rows:
            print(
                f"  {r['posted']}  {r['account_name'][:20]:<20}  "
                f"{r['description'][:35]:<35}  "
                f"{r['amount']:>10.2f}  "
                f"{(r['source_category'] or '')[:25]:<25}  "
                f"{r['category']}"
            )

        print()
        print("=== 2023 category breakdown ===")
        breakdown = conn.execute(
            """
            SELECT category, COUNT(*) as cnt, SUM(amount) as total
            FROM transactions
            WHERE posted LIKE '2023-%'
              AND source IN ('mint','empower','csv2025')
            GROUP BY category
            ORDER BY total ASC
            """
        ).fetchall()
        for r in breakdown:
            print(f"  {r['category']:<35} {r['cnt']:>5} txns  ${r['total']:>10.2f}")

    finally:
        conn.close()
